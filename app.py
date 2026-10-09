import streamlit as st
import pandas as pd
from datetime import datetime
import io
from collections import Counter
import re
import string

# =========================================================
# LISTAS DESPLEGABLES EXTRAÍDAS DE LA PLANTILLA OFICIAL
# =========================================================
LISTAS_DESPLEGABLES = {
    'Acción': ['Agregar', 'Eliminar', 'Actualizar'],
    '*Tipo Id.': ['D.N.I.', 'Pasaporte'],
    '*Incapacidad': ['SI', 'NO'],
    '*Ocupación': ['Administración del hogar. Ama de casa.-4111', 'Administrativos-1112', 'Agricultores, productores y peones de cultivos mixtos y ganaderos-6114', 'Agricultores, productores y trabajadores de cultivos (Agropecuarios, alambradores)-6111', 'Agricultores, productores y trabajadores de plantaciones de frutales, huertas, invernaderos y viveros-6113', 'Agricultores, productores y trabajadores de plantaciones de árboles y arbustos. Trabajadores forestales. (desmonte, hacheros)-6112', 'Alfareros y afines. Barro, arcilla, abrasivos.-8113', 'Alumnos de escuelas, colegios y/o complejos educativos-9111', 'Aon 1-9995', 'Aon 2-9996', 'Aon 3-9997', 'Aon 4-9998', 'Apicultores-6117', 'Artesanos de los tejidos, el cuero y materiales similares (tapiceros)-8116', 'Artesanos en madera, cestería y materiales similares-8115', 'Avicultores-6116', 'Barnizadores, lustradores y afines-7219', 'Bomberos y rescatistas-4135', 'Carboneros de carbón vegetal y afines-6121', 'Colonias de vacaciones-9112', 'Comerciantes al por mayor de bebidas y/o tabacos. Fraccionamiento, carga y descarga, transporte/distribución y venta.-5112', 'Comerciantes al por mayor de madera, papel y derivados y/o materiales para la construcción. Carga y descarga, transporte/distribución y venta.-5114', 'Comerciantes al por mayor de productos alimentarios. Acopio, carga y descarga, distribución/transporte, fraccionamiento, abastecimiento y venta.-5111', 'Comerciantes al por mayor de textiles, prendas de vestir y cuero. Distribución y venta.-5113', 'Comercios al por mayor de motores, máquinas, equipos industriales, comerciales y domésticos. Distribución y venta.-5115', 'Comisionistas y/o cobradores-2116', 'Conductores de autobuses y/o ambulancias-3113', 'Conductores de automóviles, taxis y/o camionetas (remisero, taxista)-3112', 'Conductores de motocicletas y/o bicicletas (Delivery, cadete, repartidor)-3111', 'Conductores y/o acompañantes de camiones-3114', 'Conservación de redes de agua-7612', 'Conservación de redes de gas-7611', 'Constructores y/o demoledores de edificios y/o excavadores (Construcción, zanjeo, albañiles)-7111', 'Criadores de especies acuáticas-6122', 'Criadores de ganado y trabajadores pecuarios. Tamberos. (esquiladores, cuidadores de caballos)-6115', 'Criadores y trabajadores de la cría de animales diversos y de granja-6118', 'Cristaleros. Espejos, vidrios y afines (vidrieros)-7213', 'Cuidadores y adiestradores animales domésticos-4112', 'Deportistas en la práctica de Fútbol-9138', 'Deportistas en la práctica de Hockey-9139', 'Deportistas en la práctica de artes marciales, boxeo, rugby-9144', 'Deportistas en la práctica de básquet, handball, mountain bike, airsoft, arquería-9140', 'Deportistas en la práctica de canotaje, kayak, stand up paddle, Windsurf, Kite, rapel, escalada, esquí, pesca deportiva, equitación, cabalgata-9116', 'Deportistas en la práctica de doma, jinete, jockey, pato, hipismo, sortija-9143', 'Deportistas en la práctica de fútbol, básquet, hockey, handball.-9115', 'Deportistas en la práctica de golf, paddle,  squash, tenis, voleyball, natación en piscinas, ciclismo, atletismo, gimnasia deportiva, esgrima, patinaje, newcom, trekking sin altura, running, maratón-9114', 'Directores y gerentes (Consejeros, Presidente comisión directiva)-1111', 'Electricistas (Instalación domiciliaria e industrial, alarmas)-7511', 'Empleadores / Directivos-9131', 'Empleadores / Oficios-9132', 'En tránsito vehicular-9123', 'En tránsito vehicular-9124', 'Especialistas en operaciones comerciales. Recibidores, rematadores, despachantes de aduana y Agente marítimo-2115', 'Especialistas en operaciones financieras y de administración de empresas. Contadores, productores de seguros, tasadores, agentes de bolsa y cambio. (Actuarios, administrador de propiedades, gestores, peritos)-2114', 'Expendio de combustibles. Estaciones de servicio.-5233', 'Fotógrafos (camarógrafos)-4117', 'Fumigadores terrestres-6123', 'Fumigadores, desinfecciones y otros controladores de plagas domiciliarias-4113', 'Herreros, torneros y forjadores-7311', 'Impresores y/o estampadores y operadores de máquinas de imprenta (Imprenta y litografía)-8118', 'Instaladores de material aislante (Impermeabilización de techos)-7212', 'Instaladores y reparadores de líneas eléctricas-7514', 'Instaladores y reparadores de líneas telefónicas, de tecnología de la información y comunicaciones-7516', 'Jardineros y pileteros (casas particulares)-9135', 'Joyeros, orfebres y plateros-8112', 'Lavadores y/o engrasadores de vehículos.-4114', 'Lavanderos, planchadores, tintoreros-4115', 'Limpiadores de oficinas, hoteles y otros establecimientos-4116', 'Marineros de embarcaciones de paseo-3118', 'Mecánicos montadores de instalaciones de refrigeración y climatización (Reparación de aires acondicionados en vehículos)-7215', 'Mecánicos y reparadores de máquinas agrícolas e industriales-7412', 'Mecánicos y reparadores de vehículos de motor (Taller mecánico, chapa y pintura)-7411', 'Mecánicos y reparadores en electrónica (Técnico electrónico, electromecánico, service de máquinas)-7515', 'Montadores industriales de estructuras metálicas, antenas, maquinarias. (Antenista, ascensores, instalación de máquinas frigoríficas, aire acondicionado, armado de silos y/o norias, molinos, zinguería, cartelería, toldos, carpas)-7312', 'Operadores de grúas, de aparatos elevadores y afines-3117', 'Operadores de instalaciones electrónicas, tableros y/o consolas en la industria-7712', 'Operadores de instalaciones y de máquinas en la industria (operarios de fábrica, frigoríficos, metalúrgicos, fabricación de alimentos, carpintería, clasificador de cereal, aserradero, marmolería)-7711', 'Operadores de maquinaria agrícola y forestal motorizada (cosechadoras, tractores, cargadores frontales, equipos de forraje, sembradoras, desmalezadoras, pulverizadoras, motosierra, excavadora)-3115', 'Operadores de máquinas de movimiento de tierras y afines (Grúas, máquinas viales)-3116', 'Parqueteros y colocadores de suelos (Colocación y pulidores de pisos, alfombras y similares)-7211', 'Pasantías Escolares (tareas ADMINISTRATIVAS)-9136', 'Pasantías Escolares (tareas OFICIOS / TÉCNICOS)-9137', 'Peones de carga y descarga (reparto de gas domiciliario, mudanzas, embaladores, estibadores, reparto frigorífico)-4118', 'Peones de obras públicas y mantenimiento: carreteras, presas y obras similares-7113', 'Personal de pompas fúnebres y embalsamadores-4119', 'Pintores y empapeladores (Limpiadores de vidrios, colocación cielorraso)-7216', 'Plomeros y gasistas (instaladores de tuberías de agua y gas)-7214', 'Porteros, conserjes, ordenanzas, encargados de garages y afines-4120', 'Profesionales de la enseñanza maestros, profesores, docentes y de las ciencias de la información/bibliotecarios-2113', 'Profesionales de la salud. Médicos, enfermeras, paramédicos. (Instrumentista, odontólogo, mecánico dental, radiólogo, oftalmólogo, kinesiologo, fisioterapeuta, fonoudiología, anestesista, camillero)-2112', 'Profesionales de las actividades deportivas. Instructores de educación física, entrenadores, guías de turismo, árbitros o similares.-2121', 'Profesionales de las artes y la cultura. Periodistas, traductores, escritores, locutores, músicos, artistas, actores, diseñadores.-2120', 'Profesionales de las ciencias sociales. Psicólogos, sociólogos, religiosos, Acompañante terapéutico, psicopedagoga.-2119', 'Profesionales de las ciencias y de la ingeniería. Ingenieros,  arquitectos, técnicos, bioquímicos, químicos, biólogos, físicos, geólogos (supervisor, inspectores, laboratorio, topografos y geografos, enologos, agrimensores)-2111', 'Profesionales de tecnología de la información y las comunicaciones, Desarrolladores y analistas de software, administradores de bases de datos, técnicos.-2117', 'Profesionales del derecho. Abogados, jueces o similares, escribanos, procuradores.-2118', 'Pulidores de metales y afiladores de herramientas-7314', 'Reparadores de instrumentos de precisión-8111', 'Sastres, modistos, costureros, bordadores, peleteros, sombrereros y afines-8211', 'Senderismo / Trekking Ushuaia (Residentes)-9141', 'Senderismo / Trekking Ushuaia (Turistas)-9142', 'Servicio Dom. Casas Familia M.V.-9126', 'Servicios de belleza, higiene y estética corporal (Manicura, masajista, pedicuro, peluquero, coiffeur)-4121', 'Servicios de comida y hospedaje prestados en hoteles y similares. Mucamas y cocineros.-4134', 'Servicios de comidas y bebidas prestados en eventos y similares.-4133', 'Servicios de comidas y bebidas prestados en restaurantes, bares y cafés. (cocinero, chef, mozo, barman, catering)-4132', 'Servicios de correo-4136', 'Servicios de diversión y esparcimiento en salones de baile, discotecas, salones de juegos electrónicos, billar y similares. (DJ, sonidista y armado de equipos, iluminación, decoradores de eventos)-4122', 'Servicios de jardinería y horticultura, poda, desmalezamiento y afines-4123', 'Servicios de mantenimiento general de instalaciones en establecimientos o instituciones-4124', 'Servicios de promociones y/o modelaje publicitario o artístico (promotores, modelos, encuestadores, personal en stands)-4125', 'Servicios de protección (bañeros o guardavidas en clubes o instituciones)-4126', 'Servicios de reparación de artículos personales, de los hogares y cerrajeros-4127', 'Servicios de saneamiento y similares. Recolección de residuos, limpieza, desagote de pozos (limpieza tanques de agua)-4128', 'Servicios de seguridad y vigilancia. Serenos, vigilantes, guardaparques y afines. (casero, personal de guardia)-4129', 'Servicios domésticos, cuidadores de niños, cuidadores personales, mucamas-4130', 'Servicios prestados por municipalidades o comunas-4131', 'Soldadores y oxicortadores-7315', 'Sopladores, moldeadores, laminadores, cortadores y pulidores de vidrios-8114', 'Tapiceros, colchoneros y afines-8213', 'Tareas varias por M.V.-9125', 'Transportistas-9130', 'Vendedores de aparatos fotográficos, artículos de fotografía e instrumentos de óptica.-5241', 'Vendedores de armas y cuchillería.-5226', 'Vendedores de artefactos eléctricos para iluminación.-5238', 'Vendedores de artículos de bazar y menaje. Bazares.-5235', 'Vendedores de artículos de juguetería y cotillón. Jugueterías.-5224', 'Vendedores de artículos de librería, papelería y oficina. Librerías y papelerías.-5225', 'Vendedores de artículos de tocador y cosméticos. Perfumerías.-5229', 'Vendedores de artículos para el hogar. Heladeras, lavarropas, televisores, etc.-5237', 'Vendedores de aves y huevos, animales de corral y caza y otros productos de granja.-5213', 'Vendedores de billetes de lotería y otros juegos de azar. Agencias de lotería.-5219', 'Vendedores de calzado y artículos de cuero. Zapaterías. Zapatillerías. Marroquinerías.-5221', 'Vendedores de carne y derivados. Carnicerías-5212', 'Vendedores de cámaras y cubiertas. Gomerías.-5234', 'Vendedores de flores y plantas. Distribución y venta.-5116', 'Vendedores de friambres y comidas preparadas. Rotiserías y fiambrerias.-5215', 'Vendedores de instrumentos musicales y/o discos. Disquerías. (playeros)-5223', 'Vendedores de muebles. Mueblerías.-5222', 'Vendedores de pan y demás productos de panadería. Panaderías, confiterías y/o fábrica de pastas frescas. (venta, producción, elaboración de pan y pastas)-5217', 'Vendedores de pescados y otros productos marinos, fluviales y lacustres. Pescaderías.-5214', 'Vendedores de pinturas, barnices, esmaltes, artículos de ferretería. Ferreterías y pinturerías.-5228', 'Vendedores de prendas de vestir y/o productos textiles.-5220', 'Vendedores de productos en almacenes, supermercados o autoservicios. (Repositores, cajeros).-5211', 'Vendedores de productos farmacéuticos y medicinales. Farmacias.-5227', 'Vendedores de productos medicinales para animales. Veterinarias.-5230', 'Vendedores de repuestos y accesorios para automotores.-5240', 'Vendedores de sanitarios.-5236', 'Vendedores de semillas, abonos y plaguicidas.-5231', 'Vendedores de tabaco, cigarrillos, golosinas.-5218', 'Vendedores de vehículos automotores.-5239', 'Vendedores de verduras, frutas y hortalizas frescas. Verdulerías y fruterías.-5216', 'Veterinarios de animales domésticos-2123', 'Veterinarios y vacunadores rurales-2122', 'Viajeros, excursionistas, campamentistas, turistas-9113', 'Viajes Fuera Territorio Nacional (Exterior)-9120', 'Viajes Territorio Nacional-9119', 'Vitivinicultores y/u olivicultores-6120', 'Yeseros-7218', 'Zapateros y afines-8214'],
    '*Nacionalidad': ['Afganistan ', 'Albania', 'Alemania', 'Andorra', 'Angola', 'Anguila', 'Antartida', 'Antigua y Barbuda', 'Antillas Holandesas', 'Arabia Saudita', 'Argelia', 'Argentina', 'Armenia', 'Aruba', 'Australia', 'Austria', 'Azerbaiyan', 'Bahamas', 'Bahrein', 'Banglades', 'Barbados', 'Belgica', 'Belice', 'Benin', 'Bermudas', 'Bielorrusia', 'Bolivia', 'Bosnia y Herzegovina', 'Botsuana', 'Brasil', 'Brunei', 'Bulgaria', 'Burkina Faso', 'Burundi', 'Butan', 'Cabo Verde', 'Camboya', 'Camerun', 'Canada', 'Catar', 'Chad', 'Chile', 'China', 'Chipre', 'Colombia', 'Comoras', 'Congo', 'Corea del Norte', 'Costa de Marfil', 'Costa Rica', 'Croacia', 'Cuba', 'Curazao', 'Dinamarca', 'Dominica', 'Ecuador', 'Edos. Federados de Micronesia', 'Egipto', 'El Salvador', 'Emiratos Arabes Unidos', 'Eritrea', 'Escocia', 'Eslovaquia', 'España', 'Estados Unidos de America', 'Estonia', 'Etiopia', 'Finlandia', 'Fiyi', 'Francia', 'Gabon', 'Gales', 'Gambia', 'Georgia', 'Ghana', 'Gibraltar', 'Granada', 'Grecia', 'Groenlandia', 'Guadalupe', 'Guam', 'Guatemala', 'Guayana Francesa', 'Guernsey', 'Guinea', 'Guinea Bissau', 'Guinea Ecuatorial', 'Guyana', 'Haiti', 'Hawai', 'Honduras', 'Hong Kong', 'Hungria', 'Inglaterra', 'Irak', 'Iran', 'Irlanda del Norte', 'Is. Georgia y Sandwich del Sur', 'Isla Ascension', 'Isla Bouvet', 'Isla de Man', 'Isla Navidad', 'Isla Norfolk', 'Isla Pitcairn', 'Islandia', 'Islas Caiman', 'Islas Cocos', 'Islas Cook', 'Islas Feroe', 'Islas Heard y McDonald', 'Islas Malvinas', 'Islas Marshall', 'Islas Salomon', 'Islas Svalbard y Jan Mayen', 'Islas Turcas y Caicos', 'Islas Virgenes Americanas', 'Islas Virgenes Britanicas', 'Israel', 'Italia', 'Jamaica', 'Japon', 'Jersey', 'Jordania', 'Kazajistan', 'Kenia', 'Kirguisa', 'Kiribati', 'Kuwait', 'Laos', 'Lesoto', 'Letonia', 'Libano', 'Liberia', 'Libia', 'Liechtenstein', 'Lituania', 'Luxemburgo', 'Macao', 'Macedonia', 'Madagascar', 'Malasia', 'Malaui', 'Maldivas', 'Mali', 'Malta', 'Marianas del Norte ', 'Martinica', 'Mauricio', 'Mauritania', 'Mayotte', 'Mexico', 'Moldavia', 'Monaco', 'Mongolia', 'Montenegro', 'Montserrat', 'Mozambique', 'Myanmar', 'Namibia', 'Nauru', 'Nepal', 'Nicaragua', 'Niger', 'Nigeria', 'Niue', 'Noruega', 'Nueva Caledonia', 'Nueva Zelanda', 'Oman', 'Paises Bajos', 'Pakistan', 'Palaos', 'Palestina', 'Papua Nueva Guinea', 'Paraguay', 'Peru', 'Polinesia Francesa', 'Polonia', 'Portugal', 'Puerto Rico', 'Reino de Marruecos', 'Rep. Bolivariana de Venezuela', 'Rep. Corea (Sur)', 'Rep.Democratica del Congo ', 'Republica Centroafricana', 'Republica Checa', 'Republica de Eslovenia', 'Republica de Filipinas', 'Republica de Indonesia', 'Republica de Irlanda', 'Republica de la India', 'Republica de Panama', 'Republica de Sudafrica', 'Republica Dominicana', 'Republica Oriental del Uruguay', 'Reunion', 'Ruanda', 'Rumania', 'Rusia', 'Sahara Occidental', 'Samoa', 'Samoa Americana', 'San Cristobal y Nieves', 'San Marino', 'San Pedro y Miquelon', 'San Vicente y las Granadinas', 'Santa Elena', 'Santa Lucia', 'Santo Tome y Principe', 'Senegal', 'Serbia', 'Seychelles', 'Sierra Leona', 'Singapur', 'Siria', 'Somalia', 'Sri Lanka', 'Suazilandia', 'Sudan', 'Sudan del sur', 'Suecia', 'Suiza', 'Surinam', 'Tailandia', 'Taiwan', 'Tanzania', 'Tayikistan ', 'Terr. Británico del Oc. Indico', 'Tibet', 'Timor Oriental', 'Togo', 'Tokelau', 'Tonga', 'Trinidad y Tobago', 'Tunez', 'Turkmenistan', 'Turquia', 'Tuvalu', 'Ucrania', 'Uganda', 'Uzbekistan', 'Vanuatu', 'Vaticano', 'Vietnam', 'Wallis y Futuna', 'Yemen', 'Yibuti', 'Zambia', 'Zimbabue']
}

# =========================================================
# FUNCIÓN PARA INYECTAR LAS LISTAS DESPLEGABLES EN EXCEL
# =========================================================
def inyectar_listas_desplegables(writer, df_export, workbook, sheet_name):
    worksheet_main = writer.sheets[sheet_name]
    worksheet_dv = workbook.add_worksheet('DataValidation')
    
    col_idx = 0
    for col_name, valores in LISTAS_DESPLEGABLES.items():
        worksheet_dv.write(0, col_idx, col_name)
        worksheet_dv.write_column(1, col_idx, valores)
        
        if col_name in df_export.columns:
            main_col_idx = list(df_export.columns).index(col_name)
            
            if col_idx < 26:
                letra_col_dv = string.ascii_uppercase[col_idx]
            else:
                letra_col_dv = string.ascii_uppercase[col_idx // 26 - 1] + string.ascii_uppercase[col_idx % 26]
                
            rango_formula = f"=DataValidation!${letra_col_dv}$2:${letra_col_dv}${len(valores)+1}"
            
            worksheet_main.data_validation(1, main_col_idx, 5000, main_col_idx, {
                'validate': 'list',
                'source': rango_formula
            })
            
        col_idx += 1
        
    worksheet_dv.hide()

# Configuración de la página
st.set_page_config(page_title="Asistente Operativo - AP", layout="wide")

# Fondo de pantalla y estilos tipográficos
st.markdown(
    """
    <style>
    /* IMPORTAR SORA EN SUS GROSORES: 300 (FINA), 400 (REGULAR) Y 600 (SEMI BOLD) */
    @import url('https://fonts.googleapis.com/css2?family=Sora:wght@300;400;600&display=swap');

    /* APLICAR SORA Y LINE-HEIGHT GLOBAL */
    html, body, [class*="st-"], p, h1, h2, h3, h4, h5, h6, span, label, button, li, div, small {
        font-family: 'Sora', sans-serif !important;
    }
    
    /* Texto normal Fino (300) e interlineado */
    p, li, span, div {
        font-weight: 300 !important;
        line-height: 1.6 !important;
    }

    /* Etiquetas de botones y subtextos en peso normal (400) */
    small {
        font-weight: 400 !important;
    }

    /* Títulos, subtítulos y negritas en Semi Negrita (600) */
    h1, h2, h3, h4, h5, h6, strong, b {
        font-weight: 600 !important;
    }

    /* Fondo general */
    .stApp {
        background-image: url("https://raw.githubusercontent.com/emilianopauli1985-bit/validador-ap/main/L2_Wallpaper-05.jpg");
        background-size: cover;
        background-position: center;
        background-attachment: fixed;
    }
    
    /* Paneles y tablas */
    .stAlert, [data-testid="stDataFrame"] {
        background-color: rgba(255, 255, 255, 0.95) !important;
        border-radius: 10px;
        padding: 10px;
    }

    /* Textos descriptivos de herramientas (Etiquetas de inputs) */
    label[data-testid="stWidgetLabel"] p {
        color: white !important;
        font-size: 16px !important;
        font-weight: 600 !important;
        text-shadow: 1px 1px 4px rgba(0,0,0,0.6);
    }

    /* ================= PESTAÑAS ================= */
    .stTabs [data-baseweb="tab-list"] {
        gap: 15px !important;
        background-color: transparent !important;
        flex-wrap: wrap !important;
    }

    /* Pestaña Inactiva */
    button[data-baseweb="tab"] {
        background-color: rgba(0, 0, 0, 0.6) !important; 
        border-radius: 8px !important;
        border: 1px solid rgba(255, 255, 255, 0.3) !important;
        padding: 12px 24px !important;
        margin: 0 !important;
    }
    
    /* TEXTO BLANCO EN TODAS LAS PESTAÑAS SIMULANDO AL TÍTULO */
    button[data-baseweb="tab"] div[data-testid="stMarkdownContainer"] p {
        color: white !important;
        font-size: 18px !important;
        font-weight: 600 !important; 
        text-shadow: 1px 1px 4px rgba(0,0,0,0.6) !important;
    }

    /* Pestaña Activa */
    button[data-baseweb="tab"][aria-selected="true"] {
        background-color: #900000 !important; /* Rojo oscuro elegante para resaltar el texto blanco */
        border: 2px solid white !important;
        box-shadow: 0px 4px 10px rgba(0,0,0,0.4) !important;
    }
    
    button[data-baseweb="tab"][aria-selected="true"] div[data-testid="stMarkdownContainer"] p {
        color: white !important; 
        font-weight: 600 !important; 
        text-shadow: 2px 2px 5px rgba(0,0,0,0.8) !important;
    }

    /* Ocultar barra inferior fea nativa */
    div[data-baseweb="tab-highlight"], 
    div[data-baseweb="tab-border"] {
        display: none !important;
        background-color: transparent !important;
    }
    /* ============================================ */

    /* ========================================================= */
    /* DISEÑO DEL BOTÓN DE SUBIDA DE ARCHIVOS (UPLOADER)         */
    /* ========================================================= */
    
    [data-testid="stFileUploadDropzone"] button {
        font-size: 0px !important; /* Colapsa la fuente original a 0 */
        color: transparent !important;
    }
    
    /* Oculta cualquier sub-elemento o ícono nativo de Streamlit que se cuele */
    [data-testid="stFileUploadDropzone"] button * {
        display: none !important;
    }
    
    [data-testid="stFileUploadDropzone"] button::after {
        content: "Subir archivo";
        color: #262730 !important;
        font-size: 14px !important;
        font-weight: 600 !important;
        position: absolute;
        left: 50%;
        top: 50%;
        transform: translate(-50%, -50%);
        visibility: visible !important;
        display: block !important;
    }
    
    [data-testid="stFileUploadDropzone"] small {
        font-size: 0px !important; /* Colapsa las instrucciones en inglés */
        color: transparent !important;
    }
    
    [data-testid="stFileUploadDropzone"] small::after {
        content: "Límite 200MB • Excel / CSV / PDF";
        color: rgba(49, 51, 63, 0.6) !important;
        font-size: 13px !important;
        display: block !important;
        margin-top: 2px;
    }
    </style>
    """,
    unsafe_allow_html=True
)

st.markdown('<h1 style="color: white; text-shadow: 2px 2px 5px rgba(0,0,0,0.6);">Asistente Operativo - Accidentes Personales</h1>', unsafe_allow_html=True)
st.markdown('<p style="color: white; font-size: 18px; text-shadow: 1px 1px 4px rgba(0,0,0,0.6); margin-bottom: 30px;">Seleccioná la herramienta que necesites usar hoy.</p>', unsafe_allow_html=True)

tab_armador, tab_validador, tab_pdf = st.tabs(["🪄 Armar Excel (Limpiador)", "✅ Validar Carga", "📄 PDF a Excel / Word"])

# ==========================================
# PESTAÑA 1: ARMADOR / LIMPIADOR DE DATOS CRUDOS
# ==========================================
with tab_armador:
    st.markdown('<div style="background-color: rgba(255, 255, 255, 0.95); padding: 15px; border-radius: 10px; margin-bottom: 20px;">Subí el Excel desordenado que te envió la empresa. Podés indicarle al sistema en qué fila empiezan los títulos para saltear logos o membretes.</div>', unsafe_allow_html=True)
    
    col_fila, col_espacio = st.columns([2, 3])
    with col_fila:
        fila_titulos = st.number_input("¿En qué fila del Excel están los títulos de las columnas? (Cambiá este número si hay títulos o logos arriba)", min_value=1, value=1)
    
    col1a, col2a = st.columns([2, 3])
    with col1a:
        archivo_crudo = st.file_uploader("Subí el Excel crudo del cliente", type=["xlsx", "xls", "csv"], key="uploader_armador")
    
    if archivo_crudo is not None:
        try:
            filas_a_saltar = int(fila_titulos) - 1
            if archivo_crudo.name.endswith('.csv'):
                df_crudo = pd.read_csv(archivo_crudo, skiprows=filas_a_saltar)
            else:
                df_crudo = pd.read_excel(archivo_crudo, skiprows=filas_a_saltar)
            
            df_crudo = df_crudo.dropna(how='all')
                
            columnas_disponibles = ["No incluir"] + list(df_crudo.columns)
            
            st.markdown('<div style="background-color: rgba(255, 255, 255, 0.95); padding: 15px; border-radius: 10px;">', unsafe_allow_html=True)
            st.write("### 📌 Mapeo de Columnas")
            st.write("Indicá en qué columna del Excel del cliente se encuentra cada dato:")
            
            col_map1, col_map2, col_map3 = st.columns(3)
            with col_map1:
                col_dni = st.selectbox("Columna de DNI:", options=columnas_disponibles)
            with col_map2:
                col_nombre = st.selectbox("Columna de Nombre y Apellido:", options=columnas_disponibles)
            with col_map3:
                col_fecha = st.selectbox("Columna de Fecha de Nacimiento:", options=columnas_disponibles)
                
            if st.button("Procesar y Generar Plantilla Oficial"):
                columnas_oficiales = ['Acción', '*Tipo Id.', '*Nro. Id.', '*Fecha de Nacimiento', '*Apellido', '*Nombre', '*S.A. Individual Muerte', 'S.A. Individual Inválidez', 'S.A. AMF', '*Incapacidad', '*Ocupación', '*Nacionalidad']
                df_oficial = pd.DataFrame(columns=columnas_oficiales)
                
                if col_dni != "No incluir":
                    df_oficial['*Nro. Id.'] = df_crudo[col_dni].apply(lambda x: re.sub(r'[^0-9]', '', str(x)) if str(x).lower() != 'nan' and pd.notna(x) else '')
                
                if col_nombre != "No incluir":
                    apellidos = []
                    nombres = []
                    for nombre_completo in df_crudo[col_nombre]:
                        texto = str(nombre_completo).strip()
                        if texto.lower() == 'nan' or pd.isna(nombre_completo):
                            apellidos.append("")
                            nombres.append("")
                        elif "," in texto:
                            partes = texto.split(",", 1)
                            apellidos.append(partes[0].strip().upper())
                            nombres.append(partes[1].strip().upper())
                        else:
                            partes = texto.split(" ", 1)
                            if len(partes) > 1:
                                apellidos.append(partes[0].strip().upper())
                                nombres.append(partes[1].strip().upper())
                            else:
                                apellidos.append(texto.upper())
                                nombres.append("")
                    df_oficial['*Apellido'] = apellidos
                    df_oficial['*Nombre'] = nombres
                
                if col_fecha != "No incluir":
                    fechas_limpias = []
                    for f in df_crudo[col_fecha]:
                        if pd.isna(f) or str(f).lower() == 'nan':
                            fechas_limpias.append("")
                        elif isinstance(f, (pd.Timestamp, datetime)):
                            fechas_limpias.append(f.date())
                        else:
                            f_str = str(f)
                            nums = re.sub(r'[^0-9]', '', f_str)
                            if len(nums) == 8:
                                try: fechas_limpias.append(pd.to_datetime(nums, format='%d%m%Y').date())
                                except: fechas_limpias.append(f_str)
                            else:
                                fechas_limpias.append(f_str)
                    df_oficial['*Fecha de Nacimiento'] = fechas_limpias

                df_oficial['Acción'] = 'Agregar'
                df_oficial['*Tipo Id.'] = 'D.N.I.'
                df_oficial['*Nacionalidad'] = 'Argentina'
                df_oficial['*Incapacidad'] = 'NO'
                
                st.success("✅ Plantilla generada exitosamente. Se limpiaron los datos y se aplicaron los valores por defecto y las listas desplegables.")
                st.dataframe(df_oficial)
                
                buffer_armado = io.BytesIO()
                
                with pd.ExcelWriter(buffer_armado, engine='xlsxwriter', datetime_format='dd/mm/yyyy', date_format='dd/mm/yyyy') as writer:
                    df_oficial.to_excel(writer, index=False, sheet_name="AP - Personas x Grupo")
                    workbook = writer.book
                    inyectar_listas_desplegables(writer, df_oficial, workbook, "AP - Personas x Grupo")
                
                st.download_button(
                    label="📥 Descargar Plantilla Oficial Pre-armada",
                    data=buffer_armado.getvalue(),
                    file_name="AP_Personas_Prearmado.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            st.markdown('</div>', unsafe_allow_html=True)
        except Exception as e:
            st.error(f"Error al procesar el archivo: {e}. Verificá que la fila de inicio sea la correcta.")

# ==========================================
# PESTAÑA 2: VALIDADOR ESTRICTO
# ==========================================
with tab_validador:
    col1b, col2b = st.columns([2, 3])
    with col1b:
        archivo_agente = st.file_uploader("Seleccioná el archivo Excel listo (.xlsx)", type=["xlsx"], key="uploader_validador")

    if archivo_agente is not None:
        try:
            df = pd.read_excel(archivo_agente, sheet_name="AP - Personas x Grupo")
            st.info("Procesando validaciones y corrigiendo formatos...")

            hoy = pd.Timestamp.today()
            total_fechas_mal_formato = 0
            total_fechas_corregidas = 0
            
            # CÓDIGOS PARA RANGO DE EDAD INFANTIL/ESTUDIANTIL/DEPORTIVO (1 A 80 AÑOS)
            codigos_flexibles = ['9111', '9112', '9113', '9114', '9115', '9116', '9138', '9139', '9140', '9143', '9144']
            
            errores_col = {
                'Err_TipoId': [''] * len(df), 'Err_NroId': [''] * len(df), 'Err_FechaNac': [''] * len(df),
                'Err_SAMuerte': [''] * len(df), 'Err_SAInvalidez': [''] * len(df), 'Err_SAAMF': [''] * len(df),
                'Err_Incapacidad': [''] * len(df), 'Err_Ocupacion': [''] * len(df), 'Err_Nacionalidad': [''] * len(df)
            }

            for index, row in df.iterrows():
                # --- Validar Ocupación y definir reglas de edad ---
                ocupacion_actual = str(row.get('*Ocupación', ''))
                ocup = ocupacion_actual.strip().upper()
                if pd.isna(row.get('*Ocupación')) or ocup == 'NAN' or not ocup:
                    errores_col['Err_Ocupacion'][index] = "Ocupación vacía"
                
                # Checkeamos si la ocupación tiene alguno de los códigos flexibles
                es_ocupacion_flexible = any(codigo in ocupacion_actual for codigo in codigos_flexibles)

                tipo_id = str(row.get('*Tipo Id.', '')).strip().upper()
                if pd.isna(row.get('*Tipo Id.')) or tipo_id == 'NAN' or not tipo_id:
                    errores_col['Err_TipoId'][index] = "Falta seleccionar D.N.I. o Pasaporte"
                
                nro_id = row.get('*Nro. Id.')
                nro_id_str = str(nro_id).strip()
                
                if pd.isna(nro_id) or nro_id_str == 'nan' or not nro_id_str:
                    errores_col['Err_NroId'][index] = "DNI vacío"
                else:
                    if not nro_id_str.replace('.', '').isdigit() or '.' in nro_id_str or ',' in nro_id_str:
                        errores_col['Err_NroId'][index] = "DNI con letras/puntos"
                    elif tipo_id == 'PASAPORTE':
                        if not nro_id_str.isdigit() or nro_id_str.startswith('0'):
                            errores_col['Err_NroId'][index] = "Pasaporte inicia con 0 o tiene letras"
                
                fecha_nac = row.get('*Fecha de Nacimiento')
                if pd.notna(fecha_nac):
                    es_valida = True
                    if not isinstance(fecha_nac, pd.Timestamp) and not isinstance(fecha_nac, datetime):
                        total_fechas_mal_formato += 1
                        fecha_str = str(fecha_nac)
                        numeros = re.sub(r'[^0-9]', '', fecha_str)
                        corregida = False
                        
                        if len(numeros) == 8:
                            try:
                                fecha_limpia = pd.to_datetime(numeros, format='%d%m%Y')
                                df.at[index, '*Fecha de Nacimiento'] = fecha_limpia
                                fecha_nac = fecha_limpia
                                corregida = True
                                total_fechas_corregidas += 1
                            except:
                                pass
                        
                        if not corregida:
                            errores_col['Err_FechaNac'][index] = "Fecha en formato texto irreconocible"
                            es_valida = False

                    if es_valida and (isinstance(fecha_nac, pd.Timestamp) or isinstance(fecha_nac, datetime)):
                        if fecha_nac > hoy:
                            errores_col['Err_FechaNac'][index] = "Fecha futura"
                        else:
                            edad = hoy.year - fecha_nac.year - ((hoy.month, hoy.day) < (fecha_nac.month, fecha_nac.day))
                            
                            # --- NUEVA LÓGICA DE CONTROL DE EDAD ---
                            if edad > 80:
                                errores_col['Err_FechaNac'][index] = "Mayor de 80 años"
                            elif es_ocupacion_flexible and edad < 1:
                                errores_col['Err_FechaNac'][index] = "Menor de 1 año"
                            elif not es_ocupacion_flexible and edad < 16:
                                errores_col['Err_FechaNac'][index] = "Menor de 16 años"
                                
                if pd.isna(row.get('*S.A. Individual Muerte')): errores_col['Err_SAMuerte'][index] = "S.A. Muerte vacía"
                else:
                    try: float(row.get('*S.A. Individual Muerte'))
                    except: errores_col['Err_SAMuerte'][index] = "S.A. Muerte no es número"

                if pd.isna(row.get('S.A. Individual Inválidez')): errores_col['Err_SAInvalidez'][index] = "S.A. Invalidez vacía"
                else:
                    try: float(row.get('S.A. Individual Inválidez'))
                    except: errores_col['Err_SAInvalidez'][index] = "S.A. Invalidez no es número"

                if pd.isna(row.get('S.A. AMF')): errores_col['Err_SAAMF'][index] = "S.A. AMF vacía"
                else:
                    try: float(row.get('S.A. AMF'))
                    except: errores_col['Err_SAAMF'][index] = "S.A. AMF no es número"
                        
                incap = str(row.get('*Incapacidad')).strip().upper()
                if incap not in ['SI', 'NO']:
                    errores_col['Err_Incapacidad'][index] = "Incapacidad debe ser SI o NO"
                    
                nac = str(row.get('*Nacionalidad')).strip().upper()
                if pd.isna(row.get('*Nacionalidad')) or nac == 'NAN' or not nac:
                    errores_col['Err_Nacionalidad'][index] = "Nacionalidad vacía"

            nro_ids = df['*Nro. Id.'].dropna().astype(str).tolist()
            dups = set([x for x in nro_ids if nro_ids.count(x) > 1])

            for index, row in df.iterrows():
                nro_id = str(row.get('*Nro. Id.'))
                if nro_id in dups:
                    actual = errores_col['Err_NroId'][index]
                    errores_col['Err_NroId'][index] = "DNI duplicado" if not actual else actual + " / DNI duplicado"

            lista_todos_errores = []
            for categoria, lista_err in errores_col.items():
                for error in lista_err:
                    if error != '':
                        for sub_error in error.split(" / "):
                            lista_todos_errores.append(sub_error)
            
            conteo_errores = Counter(lista_todos_errores)

            tiene_error_fila = [False] * len(df)
            mapa_nombres = {
                'Err_TipoId': 'Error: Tipo Id', 'Err_NroId': 'Error: Nro Id', 'Err_FechaNac': 'Error: Fecha Nac',
                'Err_SAMuerte': 'Error: SA Muerte', 'Err_SAInvalidez': 'Error: SA Invalidez', 'Err_SAAMF': 'Error: SA AMF',
                'Err_Incapacidad': 'Error: Incapacidad', 'Err_Ocupacion': 'Error: Ocupación', 'Err_Nacionalidad': 'Error: Nacionalidad'
            }

            columnas_con_errores = []
            for clave, nombre_col in mapa_nombres.items():
                if any(errores_col[clave]): 
                    df[nombre_col] = errores_col[clave]
                    columnas_con_errores.append(nombre_col)
                    for i, val in enumerate(errores_col[clave]):
                        if val != '': tiene_error_fila[i] = True

            df.insert(0, 'Estado Fila', ["❌ ERROR" if e else "✅ OK" for e in tiene_error_fila])
            errores_totales = sum(tiene_error_fila)
            
            if total_fechas_mal_formato > 0:
                st.success(f"🪄 **Autocorrección Inteligente:** Se detectaron {total_fechas_mal_formato} fechas mal escritas. El sistema logró corregir automáticamente {total_fechas_corregidas} de ellas.")
            
            if errores_totales == 0:
                st.success("✅ ¡Excelente! El archivo ya no tiene errores y está listo para enviar.")
            else:
                st.error(f"❌ Se encontraron {errores_totales} filas con errores que requieren intervención manual.")
                
                html_resumen = f"""
                <div style="background-color: rgba(255, 255, 255, 0.95); padding: 20px; border-radius: 10px; color: #333; margin-bottom: 20px; box-shadow: 0 4px 6px rgba(0,0,0,0.1);">
                    <h3 style="color: #d32f2f; margin-top: 0;">📊 Detalle de Errores Restantes:</h3>
                    <ul style="font-size: 16px;">
                """
                for error_texto, cantidad in conteo_errores.most_common():
                    html_resumen += f"<li><b>{cantidad}</b> x {error_texto}</li>"
                    
                html_resumen += """
                    </ul>
                    <hr style="border-top: 1px solid #ccc;">
                    <p style="margin-bottom: 0;"><b>Vista previa</b> (descargá el Excel para ver todos los detalles y usar los filtros):</p>
                </div>
                """
                st.markdown(html_resumen, unsafe_allow_html=True)
                
                st.dataframe(df[df['Estado Fila'] == "❌ ERROR"][['Estado Fila', '*Nro. Id.', '*Apellido'] + columnas_con_errores])

            df['*Fecha de Nacimiento'] = df['*Fecha de Nacimiento'].apply(lambda x: x.date() if isinstance(x, (pd.Timestamp, datetime)) else x)
            
            buffer = io.BytesIO()
            
            with pd.ExcelWriter(buffer, engine='xlsxwriter', datetime_format='dd/mm/yyyy', date_format='dd/mm/yyyy') as writer:
                df.to_excel(writer, index=False, sheet_name="AP - Personas x Grupo")
                workbook = writer.book
                worksheet = writer.sheets['AP - Personas x Grupo']
                
                inyectar_listas_desplegables(writer, df, workbook, "AP - Personas x Grupo")
                
                formato_rojo = workbook.add_format({'bg_color': '#FFC7CE', 'font_color': '#9C0006'})
                formato_verde = workbook.add_format({'bg_color': '#C6EFCE', 'font_color': '#006100'})
                formato_amarillo = workbook.add_format({'bg_color': '#FFF2CC', 'font_color': '#9C6500'})
                
                worksheet.conditional_format('A2:A5000', {'type': 'text', 'criteria': 'containing', 'value': 'ERROR', 'format': formato_rojo})
                worksheet.conditional_format('A2:A5000', {'type': 'text', 'criteria': 'containing', 'value': 'OK', 'format': formato_verde})
                worksheet.set_column(0, 0, 15) 
                
                if columnas_con_errores:
                    idx_inicio = len(df.columns) - len(columnas_con_errores)
                    worksheet.set_column(idx_inicio, len(df.columns)-1, 25, formato_amarillo)

            st.download_button(
                label="📥 Descargar Excel con Reporte de Errores" if errores_totales > 0 else "📥 Descargar Excel Validado",
                data=buffer.getvalue(),
                file_name="Control_AP_Reporte.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        except Exception as e:
            st.error(f"Error al leer el archivo. Detalle técnico: {e}")

# ==========================================
# PESTAÑA 3: CONVERTIDOR DE PDF
# ==========================================
with tab_pdf:
    st.markdown('<div style="background-color: rgba(255, 255, 255, 0.95); padding: 15px; border-radius: 10px; margin-bottom: 20px;">Subí la nómina en PDF. Si el sistema detecta una tabla estructurada, generará un <b>Excel</b>. Si solo detecta texto suelto, generará un archivo de <b>Word</b> para que te sea más fácil copiar y pegar los datos.</div>', unsafe_allow_html=True)
    
    col1c, col2c = st.columns([2, 3])
    with col1c:
        archivo_pdf = st.file_uploader("Subí el PDF del cliente", type=["pdf"], key="uploader_pdf")
    
    if archivo_pdf is not None:
        if st.button("🔍 Extraer Datos"):
            try:
                import pdfplumber
                import docx
            except ImportError:
                st.error("❌ Faltan instalar librerías. Por favor, agregá las palabras 'pdfplumber' y 'python-docx' a tu archivo requirements.txt en GitHub (una debajo de la otra) y esperá a que la aplicación se reinicie.")
            else:
                with st.spinner("Escaneando PDF... esto puede demorar unos segundos."):
                    todas_las_filas = []
                    texto_crudo = []
                    
                    try:
                        with pdfplumber.open(archivo_pdf) as pdf:
                            for page in pdf.pages:
                                tablas = page.extract_tables()
                                if tablas:
                                    for tabla in tablas:
                                        tabla_limpia = [[str(celda).strip() if celda is not None else "" for celda in fila] for fila in tabla]
                                        todas_las_filas.extend(tabla_limpia)
                                
                                texto = page.extract_text()
                                if texto:
                                    texto_crudo.extend(texto.split('\n'))
                        
                        st.markdown('<div style="background-color: rgba(255, 255, 255, 0.95); padding: 15px; border-radius: 10px;">', unsafe_allow_html=True)
                        
                        if todas_las_filas:
                            st.success("✅ ¡Se detectaron tablas estructuradas! Generando archivo Excel...")
                            df_pdf = pd.DataFrame(todas_las_filas)
                            st.write("**Vista Previa:**")
                            st.dataframe(df_pdf.head(15))
                            
                            buffer_pdf = io.BytesIO()
                            with pd.ExcelWriter(buffer_pdf, engine='xlsxwriter') as writer:
                                df_pdf.to_excel(writer, index=False, header=False, sheet_name="Datos Extraídos")
                                workbook = writer.book
                                inyectar_listas_desplegables(writer, df_pdf, workbook, "Datos Extraídos")
                            
                            st.download_button(
                                label="📥 Descargar Tabla en Excel",
                                data=buffer_pdf.getvalue(),
                                file_name="PDF_Convertido.xlsx",
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                            )
                            
                        elif texto_crudo:
                            st.warning("⚠️ El PDF no tenía formato de tabla. Se generó un archivo Word con el texto limpio para que puedas copiar y pegar.")
                            
                            doc = docx.Document()
                            doc.add_heading('Texto extraído del PDF', 0)
                            for linea in texto_crudo:
                                if linea.strip(): 
                                    doc.add_paragraph(linea.strip())
                                    
                            buffer_word = io.BytesIO()
                            doc.save(buffer_word)
                            
                            st.download_button(
                                label="📝 Descargar Texto en Word (.docx)",
                                data=buffer_word.getvalue(),
                                file_name="PDF_Texto_Extraido.docx",
                                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                            )
                        else:
                            st.error("No se detectaron textos ni tablas. Es probable que el PDF sea una foto o imagen escaneada.")
                            
                        st.markdown('</div>', unsafe_allow_html=True)
                        
                    except Exception as e:
                        st.error(f"Ocurrió un error inesperado al leer el PDF: {e}")
